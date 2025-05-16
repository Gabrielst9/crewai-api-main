from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn
import os
from datetime import datetime
from typing import Dict, List  

def criar_elemento_xml(nome: str, attrs: Dict[str, str] = None) -> OxmlElement:
    #cria elemento XML personalizado para o documento
    elemento = OxmlElement(nome)
    if attrs:
        for k, v in attrs.items():
            elemento.set(qn(k), v)
    return elemento

def configurar_documento(doc: Document) -> None:
    #define estilos base para fontes e parágrafos
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    #estilos de título
    for level in range(1, 4):
        heading = doc.styles[f'Heading {level}']
        heading.font.name = 'Calibri'
        heading.font.bold = True
        if level == 1:
            heading.font.size = Pt(14)
            heading.paragraph_format.space_before = Pt(18)
            heading.paragraph_format.space_after = Pt(12)
        elif level == 2:
            heading.font.size = Pt(12)
            heading.paragraph_format.space_before = Pt(14)
            heading.paragraph_format.space_after = Pt(10)

def criar_cabecalho(doc: Document, escola: str) -> None:
    #cria cabeçalho institucional centralizado
    section = doc.sections[0]
    header = section.header
    
    #limpa cabeçalho existente
    for paragraph in header.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    
    #adiciona texto do cabeçalho
    para = header.add_paragraph()
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = para.add_run(f"SECRETARIA DE ESTADO DA EDUCAÇÃO\nEEB {escola.upper()}")
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True

def criar_rodape(doc: Document) -> None:
    #adiciona rodapé com data de geração
    section = doc.sections[0]
    footer = section.footer
    
    #limpa rodapé existente
    for paragraph in footer.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    
    #adiciona texto do rodapé
    para = footer.add_paragraph()
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = para.add_run(f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    run.font.name = 'Calibri'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(47, 84, 150)  

def criar_tabela_informacoes(doc: Document, inputs: Dict[str, str]) -> None:
    #cria tabela formatada com dados do plano
    tabela = doc.add_table(rows=4, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = 'Light Shading Accent 1'
    
    #configura largura das colunas
    for row in tabela.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(4.2)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    #preenche com dados
    dados = [
        ("Professor:", inputs['professor']),
        ("Tema:", inputs['tema']),
        ("Série/Ano:", inputs['serie']),
        ("Duração:", inputs['duracao'])
    ]
    
    for i, (campo, valor) in enumerate(dados):
        #celula do campo
        celula = tabela.cell(i, 0)
        celula.text = campo
        celula.paragraphs[0].runs[0].font.bold = True
        celula.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        
        #celula do valor
        celula = tabela.cell(i, 1)
        celula.text = valor
        celula.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

def adicionar_secao_objetivos(doc: Document, conteudo: str) -> None:
    #adiciona lista de objetivos formatados
    doc.add_heading('OBJETIVOS DE APRENDIZAGEM', level=1)
    
    objetivos = [o.strip() for o in conteudo.split('\n') if o.strip()]
    for objetivo in objetivos[:5]:  
        para = doc.add_paragraph(style='List Bullet')
        para.add_run(objetivo)
        para.paragraph_format.left_indent = Inches(0.5)

def adicionar_secao_desenvolvimento(doc: Document) -> None:
    #adiciona seção de desenvolvimento metodológico
    doc.add_heading('DESENVOLVIMENTO METODOLÓGICO', level=1)
    
    etapas = [
        ("Introdução (15 min)", "Apresentação do tema e levantamento de conhecimentos prévios"),
        ("Desenvolvimento (50 min)", "Atividades práticas em grupo e exposição dialogada"),
        ("Conclusão (25 min)", "Sistematização dos conceitos e feedback")
    ]
    
    for titulo, descricao in etapas:
        para = doc.add_paragraph()
        para.add_run(titulo).bold = True
        para = doc.add_paragraph(descricao)
        para.paragraph_format.left_indent = Inches(0.5)
        para.paragraph_format.space_after = Pt(12)

def adicionar_secao_recursos(doc: Document, recursos: str) -> None:
    #lista recursos didáticos formatados
    doc.add_heading('RECURSOS DIDÁTICOS', level=1)
    
    recursos_lista = [r.strip() for r in recursos.split('\n') if r.strip()]
    for recurso in recursos_lista[:8]:  
        para = doc.add_paragraph(style='List Bullet')
        para.add_run(recurso)
        para.paragraph_format.left_indent = Inches(0.3)

def adicionar_secao_avaliacao(doc: Document) -> None:
    #adiciona critérios de avaliação
    doc.add_heading('AVALIAÇÃO', level=1)
    
    criterios = [
        "Participação ativa nas atividades",
        "Qualidade das produções realizadas",
        "Compreensão dos conceitos trabalhados",
        "Autoavaliação do processo de aprendizagem"
    ]
    
    for criterio in criterios:
        para = doc.add_paragraph(style='List Bullet')
        para.add_run(criterio)
        para.paragraph_format.left_indent = Inches(0.5)

def gerar_plano_aula(inputs: Dict[str, str], conteudo_plano: str, recursos: str) -> str:
    #gera documento DOCX completo
    try:
        doc = Document()
        configurar_documento(doc)
        criar_cabecalho(doc, inputs['escola'])
        criar_rodape(doc)
        
        #página de rosto
        titulo = doc.add_paragraph()
        titulo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        titulo.paragraph_format.space_before = Inches(3)
        run = titulo.add_run("PLANO DE AULA")
        run.font.size = Pt(28)
        run.font.bold = True
        
        #página de conteúdo
        doc.add_page_break()
        criar_tabela_informacoes(doc, inputs)
        doc.add_paragraph()
        
        #conteúdo pedagógico
        adicionar_secao_objetivos(doc, conteudo_plano)
        adicionar_secao_desenvolvimento(doc)
        adicionar_secao_recursos(doc, recursos)
        adicionar_secao_avaliacao(doc)
        
        #salva o documento
        os.makedirs("output", exist_ok=True)
        nome_arquivo = f"Plano_Aula_{inputs['escola']}_{inputs['tema']}.docx".replace(" ", "_")
        caminho = os.path.join("output", nome_arquivo)
        doc.save(caminho)
        
        return caminho

    except Exception as e:
        print(f"Erro ao gerar documento: {str(e)}")
        return None